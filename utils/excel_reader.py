from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK_PATH = PROJECT_ROOT / "test_data" / "OrthoLegal_Test_Data.xlsx"


def read_worksheet(
    sheet_name: str,
    workbook_path: Path = DEFAULT_WORKBOOK_PATH,
    required_columns: Iterable[str] | None = None,
) -> list[dict[str, Any]]:
    """Return every data row from a worksheet as a header-keyed dictionary."""
    if not workbook_path.exists():
        raise FileNotFoundError(f"Excel workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Worksheet '{sheet_name}' not found in workbook: {workbook_path}")

    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    if len(rows) < 2:
        raise ValueError(f"Worksheet '{sheet_name}' is empty in workbook: {workbook_path}")

    headers = list(rows[0])
    if not any(header is not None and str(header).strip() for header in headers):
        raise ValueError(f"Worksheet '{sheet_name}' has no column headers in workbook: {workbook_path}")

    header_names = [str(header).strip() if header is not None else "" for header in headers]
    required = set(required_columns or ())
    if "Execute" in required:
        required.add("Execute")
    missing_columns = [column for column in required if column not in header_names]
    if missing_columns:
        missing = ", ".join(sorted(missing_columns))
        raise ValueError(f"Worksheet '{sheet_name}' is missing required column(s): {missing}")

    return [
        {header_names[index]: value for index, value in enumerate(row) if header_names[index]}
        for row in rows[1:]
    ]


def get_test_data(
    sheet_name: str,
    action: str | None = None,
    workbook_path: Path = DEFAULT_WORKBOOK_PATH,
    required_columns: Iterable[str] | None = None,
) -> list[dict[str, Any]]:
    """Return enabled rows from a sheet, optionally filtered by Action."""
    required = set(required_columns or ())
    required.add("Execute")
    if action is not None:
        required.add("Action")

    rows = read_worksheet(sheet_name, workbook_path, required)
    enabled_rows = [row for row in rows if str(row.get("Execute", "")).strip().lower() == "yes"]
    if action is None:
        return enabled_rows

    normalized_action = action.strip().lower()
    return [
        row for row in enabled_rows
        if str(row.get("Action", "")).strip().lower() == normalized_action
    ]