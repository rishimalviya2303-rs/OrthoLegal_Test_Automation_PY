from __future__ import annotations

import ast
import os
import re
import sys
import tempfile
from copy import copy
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet


PROJECT_ROOT = Path(__file__).resolve().parents[1]
TESTS_DIR = PROJECT_ROOT / "tests"
WORKBOOK_PATH = PROJECT_ROOT / "test_data" / "Test_Execution_Commands.xlsx"
WORKSHEET_NAME = "Test Commands"
COMMAND_PREFIX = "./.venv/Scripts/python.exe -m pytest"
COMMAND_SUFFIX = "--headed -v -s"
SECRET_PATTERNS = (
    re.compile(r"password", re.I),
    re.compile(r"token", re.I),
    re.compile(r"cookie", re.I),
    re.compile(r"pin", re.I),
    re.compile(r"secret", re.I),
)


@dataclass(frozen=True)
class CommandRegistryEntry:
    test_name: str
    target: str
    command: str


def _windows_relative_path(path: Path) -> str:
    return path.relative_to(PROJECT_ROOT).as_posix()


def _read_ast(path: Path) -> ast.Module:
    return ast.parse(path.read_text(encoding="utf-8"), filename=str(path))


def _literal_keyword(call: ast.Call, keyword_name: str) -> str | None:
    for keyword in call.keywords:
        if keyword.arg == keyword_name and isinstance(keyword.value, ast.Constant):
            if isinstance(keyword.value.value, str):
                return keyword.value.value.strip()
    return None


def _test_case_name(function: ast.FunctionDef | ast.AsyncFunctionDef) -> str | None:
    for decorator in function.decorator_list:
        call = decorator if isinstance(decorator, ast.Call) else None
        if call is None:
            continue

        decorator_name = ""
        if isinstance(call.func, ast.Name):
            decorator_name = call.func.id
        elif isinstance(call.func, ast.Attribute):
            decorator_name = call.func.attr

        if decorator_name in {"test_case", "report_case"}:
            return _literal_keyword(call, "test_case_name")
    return None


def _humanize_test_name(function_name: str) -> str:
    name = function_name
    if name.startswith("test_"):
        name = name[5:]
    words = name.replace("_", " ").strip()
    return " ".join(word.upper() if word in {"pqme", "pqmer", "pqmed", "pqmes"} else word.capitalize() for word in words.split())


def _test_functions(module: ast.Module) -> list[ast.FunctionDef | ast.AsyncFunctionDef]:
    return [
        node
        for node in module.body
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and node.name.startswith("test_")
    ]


def discover_test_commands() -> list[CommandRegistryEntry]:
    entries: list[CommandRegistryEntry] = []
    for path in sorted(TESTS_DIR.rglob("test_*.py"), key=lambda item: _windows_relative_path(item).lower()):
        module = _read_ast(path)
        functions = sorted(_test_functions(module), key=lambda item: item.lineno)
        if not functions:
            continue

        relative_path = _windows_relative_path(path)
        use_node_ids = len(functions) > 1
        for function in functions:
            target = f"{relative_path}::{function.name}" if use_node_ids else relative_path
            command = f'{COMMAND_PREFIX} "{target}" {COMMAND_SUFFIX}'
            entries.append(
                CommandRegistryEntry(
                    test_name=_test_case_name(function) or _humanize_test_name(function.name),
                    target=target,
                    command=command,
                )
            )
    return entries


def _target_file(target: str) -> Path:
    relative_file = target.split("::", 1)[0]
    return PROJECT_ROOT / Path(relative_file.replace("\\", os.sep))


def validate_entries(entries: list[CommandRegistryEntry]) -> None:
    seen_targets: set[str] = set()
    discovered_functions: dict[Path, set[str]] = {}

    for entry in entries:
        if entry.target in seen_targets:
            raise ValueError(f"Duplicate test target discovered: {entry.target}")
        seen_targets.add(entry.target)

        path = _target_file(entry.target)
        try:
            path.resolve().relative_to(TESTS_DIR.resolve())
        except ValueError as exc:
            raise ValueError(f"Test target is outside tests directory: {entry.target}") from exc

        if not path.is_file():
            raise ValueError(f"Test file does not exist: {entry.target}")

        if path not in discovered_functions:
            discovered_functions[path] = {function.name for function in _test_functions(_read_ast(path))}

        if "::" in entry.target:
            function_name = entry.target.rsplit("::", 1)[1]
            if function_name not in discovered_functions[path]:
                raise ValueError(f"Test function does not exist: {entry.target}")

        if not entry.command.startswith(COMMAND_PREFIX):
            raise ValueError(f"Command has unexpected prefix: {entry.command}")
        if not entry.command.endswith(COMMAND_SUFFIX):
            raise ValueError(f"Command has unexpected suffix: {entry.command}")
        if f'"{entry.target}"' not in entry.command:
            raise ValueError(f"Command target is not quoted correctly: {entry.command}")
        if any(pattern.search(entry.command) for pattern in SECRET_PATTERNS):
            raise ValueError(f"Command may expose a secret-like value: {entry.command}")


def _format_sheet(sheet: Worksheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 62
    sheet.column_dimensions["B"].width = 132

    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for row in sheet.iter_rows():
        for cell in row:
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            alignment.vertical = "top"
            cell.alignment = alignment


def write_workbook(entries: list[CommandRegistryEntry], workbook_path: Path = WORKBOOK_PATH) -> None:
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = WORKSHEET_NAME
    sheet.append(["Test Name", "Run Command"])

    for entry in entries:
        sheet.append([entry.test_name, entry.command])

    _format_sheet(sheet)

    temp_file = tempfile.NamedTemporaryFile(
        suffix=".xlsx",
        prefix="Test_Execution_Commands_",
        dir=workbook_path.parent,
        delete=False,
    )
    temp_path = Path(temp_file.name)
    temp_file.close()

    try:
        workbook.save(temp_path)
        os.replace(temp_path, workbook_path)
    except PermissionError as exc:
        if temp_path.exists():
            temp_path.unlink()
        raise RuntimeError(
            "Unable to update Test_Execution_Commands.xlsx because the file is currently open. "
            "Close the workbook and run the sync again."
        ) from exc


def sync_test_command_registry() -> list[CommandRegistryEntry]:
    entries = discover_test_commands()
    validate_entries(entries)
    write_workbook(entries)
    return entries


def validate_workbook(workbook_path: Path = WORKBOOK_PATH) -> list[CommandRegistryEntry]:
    entries = discover_test_commands()
    validate_entries(entries)

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if WORKSHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"Worksheet not found: {WORKSHEET_NAME}")
        sheet = workbook[WORKSHEET_NAME]
        headers = [sheet["A1"].value, sheet["B1"].value]
        if headers != ["Test Name", "Run Command"]:
            raise ValueError(f"Unexpected workbook headers: {headers}")
        commands = [row[1] for row in sheet.iter_rows(min_row=2, values_only=True) if row[1]]
        expected_commands = [entry.command for entry in entries]
        if commands != expected_commands:
            raise ValueError("Workbook commands do not match current discovered tests.")
        if len(commands) != len(set(commands)):
            raise ValueError("Workbook contains duplicate commands.")
    finally:
        workbook.close()

    return entries


def main() -> int:
    try:
        entries = sync_test_command_registry()
        validate_workbook()
    except Exception as exc:
        print(f"Test command registry sync failed: {exc}", file=sys.stderr)
        return 1

    print(f"Updated {WORKBOOK_PATH}")
    print(f"Worksheet: {WORKSHEET_NAME}")
    print(f"Rows written: {len(entries)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
