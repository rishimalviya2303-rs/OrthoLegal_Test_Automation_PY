import argparse
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.excel_reader import PROJECT_ROOT


ENCOUNTER_WORKBOOK_PATH = PROJECT_ROOT / "test_data" / "Encounter_Test_Data.xlsx"
ENCOUNTER_SHEET_NAME = "Encounter_Test_Data"
PQME_ENCOUNTER_WORKBOOK_PATH = PROJECT_ROOT / "test_data" / "PQME_Encounter_Test_Data.xlsx"
PQME_ENCOUNTER_SHEET_NAME = ENCOUNTER_SHEET_NAME
PQMER_ENCOUNTER_WORKBOOK_PATH = PROJECT_ROOT / "test_data" / "PQMER_Encounter_Test_Data.xlsx"
PQMER_ENCOUNTER_SHEET_NAME = ENCOUNTER_SHEET_NAME
PQMES_ENCOUNTER_WORKBOOK_PATH = PROJECT_ROOT / "test_data" / "PQMES_Encounter_Test_Data.xlsx"
PQMES_ENCOUNTER_SHEET_NAME = "PQMES_Encounter_Test_Data"
REFLEX_PQMES_ENCOUNTER_WORKBOOK_PATH = PROJECT_ROOT / "test_data" / "Reflex_PQMES_Encounter_Test_Data.xlsx"
REFLEX_PQMES_ENCOUNTER_SHEET_NAME = "Reflex_PQMES_Encounter_Test_Dat"
PQMED_ENCOUNTER_WORKBOOK_PATH = PROJECT_ROOT / "test_data" / "PQMED_Encounter_Test_Data.xlsx"
PQMED_ENCOUNTER_SHEET_NAME = "PQMED_Encounter_Test_Data"
AME_ENCOUNTER_WORKBOOK_PATH = PROJECT_ROOT / "test_data" / "AME_Encounter_Test_Data.xlsx"
AME_ENCOUNTER_SHEET_NAME = "AME_Encounter_Test_Data"
AMER_ENCOUNTER_WORKBOOK_PATH = PROJECT_ROOT / "test_data" / "AMER_Encounter_Test_Data.xlsx"
AMER_ENCOUNTER_SHEET_NAME = "AMER_Encounter_Test_Data"
AMES_ENCOUNTER_WORKBOOK_PATH = PROJECT_ROOT / "test_data" / "AMES_Encounter_Test_Data.xlsx"
AMES_ENCOUNTER_SHEET_NAME = "AMES_Encounter_Test_Data"
REFLEX_AMES_ENCOUNTER_WORKBOOK_PATH = PROJECT_ROOT / "test_data" / "Reflex_AMES_Encounter_Test_Data.xlsx"
REFLEX_AMES_ENCOUNTER_SHEET_NAME = "Reflex_AMES_Encounter_Test_Data"
AMED_ENCOUNTER_WORKBOOK_PATH = PROJECT_ROOT / "test_data" / "AMED_Encounter_Test_Data.xlsx"
AMED_ENCOUNTER_SHEET_NAME = "AMED_Encounter_Test_Data"

PQME_ENCOUNTER_DETAILS_FIELDS = (
    "Physician",
    "Medical Evaluation Type",
    "Encounter Type",
    "Panel Number",
    "Date of Appointment Call",
    "Appointment Date and Time",
    "Requesting Party",
    "Examination Address",
    "City",
    "State",
    "Zip",
    "Evaluation Type",
)

PQMES_ENCOUNTER_DETAILS_FIELDS = (
    "Physician",
    "Medical Evaluation Type",
    "Encounter Type",
    "Panel Number",
    "Date of Cover Letter",
    "Date of Declaration",
    "Requesting Party",
)

REFLEX_PQMES_ENCOUNTER_DETAILS_FIELDS = (
    "Physician",
    "Medical Evaluation Type",
    "Encounter Type",
    "Panel Number",
    "Date of Receipt",
    "Requesting Party",
)

REFLEX_AMES_ENCOUNTER_DETAILS_FIELDS = (
    "Physician",
    "Medical Evaluation Type",
    "Encounter Type",
    "Date of Receipt",
    "Requesting Party",
)

PQMED_ENCOUNTER_DETAILS_FIELDS = (
    "Physician",
    "Medical Evaluation Type",
    "Encounter Type",
    "Panel Number",
    "Date of Appointment Call",
    "Appointment Date and Time",
    "Requesting Party",
    "Address Line 1",
    "Address Line 2",
    "City",
    "State",
    "Zip",
)

AMED_ENCOUNTER_DETAILS_FIELDS = (
    "Physician",
    "Medical Evaluation Type",
    "Encounter Type",
    "Date of Appointment Call",
    "Appointment Date and Time",
    "Requesting Party",
    "Address Line 1",
    "Address Line 2",
    "City",
    "State",
    "Zip",
)

AME_ENCOUNTER_DETAILS_FIELDS = (
    "Physician",
    "Medical Evaluation Type",
    "Encounter Type",
    "Date of Appointment Call",
    "Appointment Date and Time",
    "Requesting Party",
    "Examination Address",
    "City",
    "State",
    "Zip",
    "Evaluation Type",
)

AMES_ENCOUNTER_DETAILS_FIELDS = (
    "Physician",
    "Medical Evaluation Type",
    "Encounter Type",
    "Date of Cover Letter",
    "Date of Declaration",
    "Requesting Party",
)

INTERPRETER_DETAILS_FIELDS = (
    "Language",
    "Vendor Name",
    "Vendor Phone",
)

ENCOUNTER_SECTIONS: tuple[tuple[str, tuple[str, ...]], ...] = (
    (
        "Applicant Details",
        (
            "First Name",
            "Middle Name",
            "Last Name",
            "Date of Birth",
            "Gender",
            "SSN",
            "Address Line 1",
            "Address Line 2",
            "City",
            "State",
            "Zip",
            "Home Phone",
            "Cell Phone",
            "Work Phone",
            "Email",
            "Job Title",
        ),
    ),
    (
        "Injury Details",
        (
            "Date of Injury",
            "Injured Body Parts",
            "WCAB Number",
            "Claim Number",
        ),
    ),
    (
        "Encounter Details",
        PQME_ENCOUNTER_DETAILS_FIELDS,
    ),
    (
        "Upload Documents",
        (
            "Upload Strike Panel",
            "Upload Attorney Demos",
        ),
    ),
    (
        "Applicant Attorney Details",
        (
            "Firm",
            "First Name",
            "Middle Name",
            "Last Name",
            "Address Line 1",
            "Address Line 2",
            "City",
            "State",
            "Zip",
            "Phone",
            "Fax",
            "Email",
            "Assistant Email",
        ),
    ),
    (
        "Employer Details",
        (
            "Employer",
            "Address Line 1",
            "Address Line 2",
            "City",
            "State",
            "Zip",
            "Phone",
            "Fax",
            "Email",
        ),
    ),
    (
        "Defense Attorney Details",
        (
            "Firm",
            "First Name",
            "Middle Name",
            "Last Name",
            "Address Line 1",
            "Address Line 2",
            "City",
            "State",
            "Zip",
            "Phone",
            "Fax",
            "Email",
            "Assistant Email",
        ),
    ),
    (
        "Insurance Company Details",
        (
            "Insurance Company",
            "Payer ID",
            "Insurance Company Phone",
            "Adjuster First Name",
            "Adjuster Middle Name",
            "Adjuster Last Name",
            "Address Line 1",
            "Address Line 2",
            "City",
            "State",
            "Zip",
            "Adjuster Phone",
            "Adjuster Fax",
            "Adjuster Email",
            "Adjuster Assistant Email",
        ),
    ),
)

DATE_FIELDS = {
    ("Applicant Details", "Date of Birth"),
    ("Injury Details", "Date of Injury"),
    ("Encounter Details", "Date of Appointment Call"),
    ("Encounter Details", "Date of Cover Letter"),
    ("Encounter Details", "Date of Declaration"),
    ("Encounter Details", "Date of Receipt"),
}
DATETIME_FIELDS = {
    ("Encounter Details", "Appointment Date and Time"),
}


def encounter_sections(
    encounter_details_fields: tuple[str, ...] = PQME_ENCOUNTER_DETAILS_FIELDS,
    include_interpreter_details: bool = False,
) -> tuple[tuple[str, tuple[str, ...]], ...]:
    """Return Encounter sections with a configurable Encounter Details field list.

    When *include_interpreter_details* is ``True`` the ``Interpreter Details``
    section is inserted immediately after ``Upload Documents``.
    """
    sections: list[tuple[str, tuple[str, ...]]] = []
    for section, fields in ENCOUNTER_SECTIONS:
        if section == "Encounter Details":
            sections.append((section, encounter_details_fields))
        else:
            sections.append((section, fields))
        if include_interpreter_details and section == "Upload Documents":
            sections.append(("Interpreter Details", INTERPRETER_DETAILS_FIELDS))
    return tuple(sections)


def encounter_columns(
    encounter_details_fields: tuple[str, ...] = PQME_ENCOUNTER_DETAILS_FIELDS,
    include_interpreter_details: bool = False,
) -> list[tuple[str, str]]:
    """Return Encounter columns as ordered (section, field) pairs."""
    return [
        (section, field)
        for section, fields in encounter_sections(
            encounter_details_fields, include_interpreter_details=include_interpreter_details
        )
        for field in fields
    ]


def create_encounter_test_data_workbook(
    workbook_path: Path = ENCOUNTER_WORKBOOK_PATH,
    sheet_name: str = ENCOUNTER_SHEET_NAME,
    encounter_details_fields: tuple[str, ...] = PQME_ENCOUNTER_DETAILS_FIELDS,
    include_interpreter_details: bool = False,
) -> Path:
    """Create the blank Encounter Creation test-data workbook."""
    workbook_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.freeze_panes = "A3"

    section_fill = PatternFill("solid", fgColor="1F4E78")
    field_fill = PatternFill("solid", fgColor="D9EAF7")
    section_font = Font(bold=True, color="FFFFFF")
    field_font = Font(bold=True, color="000000")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    column_index = 1
    for section, fields in encounter_sections(
        encounter_details_fields, include_interpreter_details=include_interpreter_details
    ):
        start_column = column_index
        for field in fields:
            section_cell = worksheet.cell(row=1, column=column_index, value=section)
            field_cell = worksheet.cell(row=2, column=column_index, value=field)
            section_cell.fill = section_fill
            section_cell.font = section_font
            section_cell.alignment = center
            field_cell.fill = field_fill
            field_cell.font = field_font
            field_cell.alignment = center
            worksheet.column_dimensions[get_column_letter(column_index)].width = max(16, min(28, len(field) + 4))
            column_index += 1

        end_column = column_index - 1
        if end_column > start_column:
            worksheet.merge_cells(
                start_row=1,
                start_column=start_column,
                end_row=1,
                end_column=end_column,
            )

    worksheet.row_dimensions[1].height = 24
    worksheet.row_dimensions[2].height = 36

    text_format = "@"
    for row in range(3, 203):
        for col in range(1, column_index):
            worksheet.cell(row=row, column=col).number_format = text_format

    workbook.save(workbook_path)
    workbook.close()
    return workbook_path


def read_encounter_test_data(
    row_number: int = 3,
    workbook_path: Path = ENCOUNTER_WORKBOOK_PATH,
    sheet_name: str = ENCOUNTER_SHEET_NAME,
) -> dict[str, dict[str, Any]]:
    """Read one Encounter data row grouped by section and exact field name."""
    if row_number < 3:
        raise ValueError("Encounter data rows start at Excel row 3.")
    if not workbook_path.exists():
        raise FileNotFoundError(f"Excel workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"Worksheet '{sheet_name}' not found in workbook: {workbook_path}")

    worksheet = workbook[sheet_name]
    columns = _read_header_columns(worksheet)
    data: dict[str, dict[str, Any]] = {}
    for column_index, section, field in columns:
        value = worksheet.cell(row=row_number, column=column_index).value
        if (section, field) in DATE_FIELDS:
            value = format_encounter_date_value(value, section=section, field=field)
        elif (section, field) in DATETIME_FIELDS:
            value = format_encounter_datetime_value(value, section=section, field=field)
        data.setdefault(section, {})[field] = value

    workbook.close()
    return data


def format_encounter_date_value(value: Any, *, section: str, field: str) -> str | None:
    """Return Encounter date input values in MM/DD/YYYY format."""
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.strftime("%m/%d/%Y")

    if isinstance(value, date):
        return value.strftime("%m/%d/%Y")

    text = str(value).strip()
    if not text:
        return None

    date_part = text.split()[0]
    for pattern in ("%m/%d/%Y", "%m-%d-%Y", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            parsed = datetime.strptime(date_part, pattern)
            return parsed.strftime("%m/%d/%Y")
        except ValueError:
            continue

    raise ValueError(
        f"Encounter date field '{section} -> {field}' must be a valid date "
        "that can be formatted as MM/DD/YYYY."
    )


def format_encounter_datetime_value(value: Any, *, section: str, field: str) -> str | None:
    """Return Encounter date-time input values with MM/DD/YYYY date formatting."""
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.strftime("%m/%d/%Y %I:%M %p")

    text = str(value).strip()
    if not text:
        return None

    for pattern in (
        "%m/%d/%Y %I:%M %p",
        "%m-%d-%Y %I:%M %p",
        "%m/%d/%Y %H:%M",
        "%m-%d-%Y %H:%M",
        "%Y-%m-%d %H:%M",
    ):
        try:
            parsed = datetime.strptime(text, pattern)
            return parsed.strftime("%m/%d/%Y %I:%M %p")
        except ValueError:
            continue

    date_only = format_encounter_date_value(value, section=section, field=field)
    if date_only:
        return date_only
    return None


def validate_encounter_workbook(
    workbook_path: Path = ENCOUNTER_WORKBOOK_PATH,
    sheet_name: str = ENCOUNTER_SHEET_NAME,
    encounter_details_fields: tuple[str, ...] = PQME_ENCOUNTER_DETAILS_FIELDS,
    include_interpreter_details: bool = False,
) -> None:
    """Validate workbook headers against the Encounter Creation schema."""
    if not workbook_path.exists():
        raise FileNotFoundError(f"Excel workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if workbook.sheetnames != [sheet_name]:
        workbook.close()
        raise ValueError(f"Workbook must contain only the '{sheet_name}' worksheet.")

    worksheet = workbook[sheet_name]
    actual_columns = [(section, field) for _, section, field in _read_header_columns(worksheet)]
    expected_columns = encounter_columns(
        encounter_details_fields, include_interpreter_details=include_interpreter_details
    )
    workbook.close()

    if actual_columns != expected_columns:
        raise ValueError("Encounter workbook headers do not match the expected section and field order.")


def _read_header_columns(worksheet: Any) -> list[tuple[int, str, str]]:
    columns: list[tuple[int, str, str]] = []
    current_section = ""

    for column_index in range(1, worksheet.max_column + 1):
        section_value = worksheet.cell(row=1, column=column_index).value
        field_value = worksheet.cell(row=2, column=column_index).value
        if section_value is not None and str(section_value).strip():
            current_section = str(section_value).strip()
        if field_value is None or not str(field_value).strip():
            continue
        if not current_section:
            raise ValueError(f"Missing section header for column {column_index}.")
        columns.append((column_index, current_section, str(field_value).strip()))

    if not columns:
        raise ValueError("Encounter worksheet has no two-level headers.")
    return columns


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Create Encounter test-data workbooks.")
    parser.add_argument(
        "--template",
        choices=("pqme", "pqmer", "pqmes", "pqmed", "reflex-pqmes", "ame", "amer", "ames", "reflex-ames", "amed"),
        default="pqme",
        help="Encounter workbook template to generate.",
    )
    args = parser.parse_args()

    if args.template == "pqmes":
        output_path = create_encounter_test_data_workbook(
            workbook_path=PQMES_ENCOUNTER_WORKBOOK_PATH,
            sheet_name=PQMES_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=PQMES_ENCOUNTER_DETAILS_FIELDS,
        )
        validate_encounter_workbook(
            output_path,
            sheet_name=PQMES_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=PQMES_ENCOUNTER_DETAILS_FIELDS,
        )
    elif args.template == "reflex-pqmes":
        output_path = create_encounter_test_data_workbook(
            workbook_path=REFLEX_PQMES_ENCOUNTER_WORKBOOK_PATH,
            sheet_name=REFLEX_PQMES_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=REFLEX_PQMES_ENCOUNTER_DETAILS_FIELDS,
        )
        validate_encounter_workbook(
            output_path,
            sheet_name=REFLEX_PQMES_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=REFLEX_PQMES_ENCOUNTER_DETAILS_FIELDS,
        )
    elif args.template == "pqmed":
        output_path = create_encounter_test_data_workbook(
            workbook_path=PQMED_ENCOUNTER_WORKBOOK_PATH,
            sheet_name=PQMED_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=PQMED_ENCOUNTER_DETAILS_FIELDS,
            include_interpreter_details=True,
        )
        validate_encounter_workbook(
            output_path,
            sheet_name=PQMED_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=PQMED_ENCOUNTER_DETAILS_FIELDS,
            include_interpreter_details=True,
        )
    elif args.template == "pqmer":
        output_path = create_encounter_test_data_workbook(
            workbook_path=PQMER_ENCOUNTER_WORKBOOK_PATH,
            sheet_name=PQMER_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=PQME_ENCOUNTER_DETAILS_FIELDS,
            include_interpreter_details=True,
        )
        validate_encounter_workbook(
            output_path,
            sheet_name=PQMER_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=PQME_ENCOUNTER_DETAILS_FIELDS,
            include_interpreter_details=True,
        )
    elif args.template == "ame":
        output_path = create_encounter_test_data_workbook(
            workbook_path=AME_ENCOUNTER_WORKBOOK_PATH,
            sheet_name=AME_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=AME_ENCOUNTER_DETAILS_FIELDS,
            include_interpreter_details=True,
        )
        validate_encounter_workbook(
            output_path,
            sheet_name=AME_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=AME_ENCOUNTER_DETAILS_FIELDS,
            include_interpreter_details=True,
        )
    elif args.template == "amer":
        output_path = create_encounter_test_data_workbook(
            workbook_path=AMER_ENCOUNTER_WORKBOOK_PATH,
            sheet_name=AMER_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=AME_ENCOUNTER_DETAILS_FIELDS,
            include_interpreter_details=True,
        )
        validate_encounter_workbook(
            output_path,
            sheet_name=AMER_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=AME_ENCOUNTER_DETAILS_FIELDS,
            include_interpreter_details=True,
        )
    elif args.template == "ames":
        output_path = create_encounter_test_data_workbook(
            workbook_path=AMES_ENCOUNTER_WORKBOOK_PATH,
            sheet_name=AMES_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=AMES_ENCOUNTER_DETAILS_FIELDS,
        )
        validate_encounter_workbook(
            output_path,
            sheet_name=AMES_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=AMES_ENCOUNTER_DETAILS_FIELDS,
        )
    elif args.template == "reflex-ames":
        output_path = create_encounter_test_data_workbook(
            workbook_path=REFLEX_AMES_ENCOUNTER_WORKBOOK_PATH,
            sheet_name=REFLEX_AMES_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=REFLEX_AMES_ENCOUNTER_DETAILS_FIELDS,
        )
        validate_encounter_workbook(
            output_path,
            sheet_name=REFLEX_AMES_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=REFLEX_AMES_ENCOUNTER_DETAILS_FIELDS,
        )
    elif args.template == "amed":
        output_path = create_encounter_test_data_workbook(
            workbook_path=AMED_ENCOUNTER_WORKBOOK_PATH,
            sheet_name=AMED_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=AMED_ENCOUNTER_DETAILS_FIELDS,
            include_interpreter_details=True,
        )
        validate_encounter_workbook(
            output_path,
            sheet_name=AMED_ENCOUNTER_SHEET_NAME,
            encounter_details_fields=AMED_ENCOUNTER_DETAILS_FIELDS,
            include_interpreter_details=True,
        )
    else:
        output_path = create_encounter_test_data_workbook(
            include_interpreter_details=True,
        )
        validate_encounter_workbook(output_path, include_interpreter_details=True)
    print(f"Created {output_path}")
