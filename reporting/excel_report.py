from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


REPORT_DIRECTORY = Path(__file__).resolve().parents[1] / "reports"
HEADERS = [
    "Test Case ID",
    "Module",
    "Test Case Title",
    "Test Type",
    "Preconditions",
    "Test Steps",
    "Test Data",
    "Expected Result",
    "Actual Result",
    "Status",
    "Execution Date/Time",
]
IST = ZoneInfo("Asia/Kolkata")
_report_path: Path | None = None
_execution_timestamp: str | None = None


def initialize_report() -> Path:
    global _execution_timestamp, _report_path

    execution_time = datetime.now(IST)
    _execution_timestamp = execution_time.strftime("%d %b, %Y %I:%M %p IST")
    filename_timestamp = execution_time.strftime("%d_%b_%Y_%I-%M-%S_%p_IST")
    REPORT_DIRECTORY.mkdir(parents=True, exist_ok=True)
    base_path = REPORT_DIRECTORY / f"OrthoLegal_Automation_Test_Report_{filename_timestamp}.xlsx"
    _report_path = base_path
    suffix = 1
    while _report_path.exists():
        _report_path = REPORT_DIRECTORY / f"OrthoLegal_Automation_Test_Report_{filename_timestamp}_{suffix}.xlsx"
        suffix += 1
    return _report_path


def record_test_case(test_case: dict[str, Any], status: str, actual_result: str) -> None:
    if _report_path is None or _execution_timestamp is None:
        initialize_report()

    assert _report_path is not None
    assert _execution_timestamp is not None
    workbook = load_workbook(_report_path) if _report_path.exists() else Workbook()
    worksheet = workbook[test_case["module"]] if test_case["module"] in workbook.sheetnames else workbook.create_sheet(test_case["module"])
    _remove_empty_default_sheet(workbook)

    _ensure_headers(worksheet, _execution_timestamp)
    row_number = _find_test_case_row(worksheet, test_case["test_case_id"])
    if row_number is None:
        row_number = max(worksheet.max_row + 1, 4)

    values = [
        test_case["test_case_id"],
        test_case["module"],
        test_case["title"],
        test_case["test_type"],
        test_case["preconditions"],
        test_case["steps"],
        test_case["test_data"],
        test_case["expected_result"],
        actual_result,
        status,
        _execution_timestamp,
    ]
    for column_number, value in enumerate(values, start=1):
        cell = worksheet.cell(row=row_number, column=column_number, value=value)
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    _format_worksheet(worksheet)
    workbook.save(_report_path)


def _remove_empty_default_sheet(workbook: Any) -> None:
    default_sheet = workbook["Sheet"] if "Sheet" in workbook.sheetnames else None
    if default_sheet is not None and default_sheet.max_row == 1 and default_sheet.max_column == 1 and default_sheet["A1"].value is None and len(workbook.sheetnames) > 1:
        workbook.remove(default_sheet)


def _ensure_headers(worksheet: Any, execution_timestamp: str) -> None:
    worksheet["A1"] = "Test Execution:"
    worksheet["B1"] = execution_timestamp
    worksheet["A1"].font = Font(bold=True)
    worksheet["A1"].alignment = Alignment(vertical="top", wrap_text=True)
    worksheet["B1"].alignment = Alignment(vertical="top", wrap_text=True)

    if worksheet.max_row < 3 or all(cell.value is None for cell in worksheet[3]):
        for column_number, header in enumerate(HEADERS, start=1):
            worksheet.cell(row=3, column=column_number, value=header)
    elif [cell.value for cell in worksheet[3]] != HEADERS:
        for column_number, header in enumerate(HEADERS, start=1):
            worksheet.cell(row=3, column=column_number, value=header)

    for cell in worksheet[3]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    worksheet.freeze_panes = "A4"


def _find_test_case_row(worksheet: Any, test_case_id: str) -> int | None:
    for row_number in range(4, worksheet.max_row + 1):
        if worksheet.cell(row=row_number, column=1).value == test_case_id:
            return row_number
    return None


def _format_worksheet(worksheet: Any) -> None:
    widths = [18, 16, 58, 14, 34, 55, 34, 55, 55, 12, 22]
    for column_number, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column_number)].width = width
    for row in worksheet.iter_rows(min_row=4):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
