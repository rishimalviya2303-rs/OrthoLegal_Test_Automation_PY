from datetime import datetime

from openpyxl import load_workbook

from utils.encounter_test_data import (
    DATE_FIELDS,
    ENCOUNTER_SECTIONS,
    ENCOUNTER_SHEET_NAME,
    INTERPRETER_DETAILS_FIELDS,
    PQMES_ENCOUNTER_DETAILS_FIELDS,
    PQMES_ENCOUNTER_SHEET_NAME,
    create_encounter_test_data_workbook,
    encounter_columns,
    format_encounter_date_value,
    format_encounter_datetime_value,
    read_encounter_test_data,
    validate_encounter_workbook,
)


def test_create_encounter_test_data_workbook_uses_two_level_headers(tmp_path) -> None:
    workbook_path = tmp_path / "Encounter_Test_Data.xlsx"

    create_encounter_test_data_workbook(workbook_path)
    validate_encounter_workbook(workbook_path)

    workbook = load_workbook(workbook_path, data_only=True)
    worksheet = workbook[ENCOUNTER_SHEET_NAME]

    actual_columns = []
    current_section = None
    for column_index in range(1, worksheet.max_column + 1):
        section = worksheet.cell(row=1, column=column_index).value
        field = worksheet.cell(row=2, column=column_index).value
        current_section = section or current_section
        actual_columns.append((current_section, field))

    assert workbook.sheetnames == [ENCOUNTER_SHEET_NAME]
    assert worksheet.freeze_panes == "A3"
    assert actual_columns == encounter_columns()
    assert worksheet.max_column == 84
    assert all(
        worksheet.cell(row=3, column=column_index).value is None
        for column_index in range(1, worksheet.max_column + 1)
    )
    workbook.close()


def test_read_encounter_test_data_groups_repeated_field_names_by_section(tmp_path) -> None:
    workbook_path = tmp_path / "Encounter_Test_Data.xlsx"
    create_encounter_test_data_workbook(workbook_path)

    workbook = load_workbook(workbook_path)
    worksheet = workbook[ENCOUNTER_SHEET_NAME]
    column_lookup = {
        (section, field): index
        for index, (section, field) in enumerate(encounter_columns(), start=1)
    }
    worksheet.cell(row=3, column=column_lookup[("Applicant Details", "First Name")], value="Applicant")
    worksheet.cell(row=3, column=column_lookup[("Applicant Attorney Details", "First Name")], value="Attorney")
    worksheet.cell(row=3, column=column_lookup[("Defense Attorney Details", "First Name")], value="Defense")
    worksheet.cell(row=3, column=column_lookup[("Insurance Company Details", "Adjuster Email")], value="adjuster@example.test")
    workbook.save(workbook_path)
    workbook.close()

    encounter_data = read_encounter_test_data(workbook_path=workbook_path)

    assert list(encounter_data) == [section for section, _ in ENCOUNTER_SECTIONS]
    assert encounter_data["Applicant Details"]["First Name"] == "Applicant"
    assert encounter_data["Applicant Attorney Details"]["First Name"] == "Attorney"
    assert encounter_data["Defense Attorney Details"]["First Name"] == "Defense"
    assert encounter_data["Insurance Company Details"]["Adjuster Email"] == "adjuster@example.test"
    assert encounter_data["Employer Details"]["Fax"] is None


def test_read_encounter_test_data_supports_custom_sheet_name(tmp_path) -> None:
    workbook_path = tmp_path / "PQME_Encounter_Test_Data.xlsx"
    create_encounter_test_data_workbook(
        workbook_path,
        sheet_name=PQMES_ENCOUNTER_SHEET_NAME,
        encounter_details_fields=PQMES_ENCOUNTER_DETAILS_FIELDS,
    )
    validate_encounter_workbook(
        workbook_path,
        sheet_name=PQMES_ENCOUNTER_SHEET_NAME,
        encounter_details_fields=PQMES_ENCOUNTER_DETAILS_FIELDS,
    )

    workbook = load_workbook(workbook_path)
    worksheet = workbook[PQMES_ENCOUNTER_SHEET_NAME]
    column_lookup = {
        (section, field): index
        for index, (section, field) in enumerate(encounter_columns(PQMES_ENCOUNTER_DETAILS_FIELDS), start=1)
    }
    worksheet.cell(row=3, column=column_lookup[("Encounter Details", "Encounter Type")], value="PQMES")
    workbook.save(workbook_path)
    workbook.close()

    encounter_data = read_encounter_test_data(
        workbook_path=workbook_path,
        sheet_name=PQMES_ENCOUNTER_SHEET_NAME,
    )

    assert encounter_data["Encounter Details"]["Encounter Type"] == "PQMES"


def test_read_encounter_test_data_formats_date_fields_as_mm_dd_yyyy(tmp_path) -> None:
    workbook_path = tmp_path / "Encounter_Test_Data.xlsx"
    create_encounter_test_data_workbook(workbook_path)

    workbook = load_workbook(workbook_path)
    worksheet = workbook[ENCOUNTER_SHEET_NAME]
    column_lookup = {
        (section, field): index
        for index, (section, field) in enumerate(encounter_columns(), start=1)
    }
    worksheet.cell(row=3, column=column_lookup[("Applicant Details", "Date of Birth")], value="09-06-2026")
    worksheet.cell(row=3, column=column_lookup[("Injury Details", "Date of Injury")], value="09/07/2026")
    worksheet.cell(row=3, column=column_lookup[("Encounter Details", "Date of Appointment Call")], value=datetime(2026, 9, 8))
    worksheet.cell(row=3, column=column_lookup[("Encounter Details", "Appointment Date and Time")], value="2026-09-09 10:30")
    workbook.save(workbook_path)
    workbook.close()

    encounter_data = read_encounter_test_data(workbook_path=workbook_path)

    assert encounter_data["Applicant Details"]["Date of Birth"] == "09/06/2026"
    assert encounter_data["Injury Details"]["Date of Injury"] == "09/07/2026"
    assert encounter_data["Encounter Details"]["Date of Appointment Call"] == "09/08/2026"
    assert encounter_data["Encounter Details"]["Appointment Date and Time"] == "09/09/2026 10:30 AM"
    assert DATE_FIELDS == {
        ("Applicant Details", "Date of Birth"),
        ("Injury Details", "Date of Injury"),
        ("Encounter Details", "Date of Appointment Call"),
        ("Encounter Details", "Date of Cover Letter"),
        ("Encounter Details", "Date of Declaration"),
        ("Encounter Details", "Date of Receipt"),
    }


def test_format_encounter_datetime_value_requires_mm_dd_yyyy_date_output() -> None:
    assert (
        format_encounter_datetime_value(
            "09-06-2026 14:05",
            section="Encounter Details",
            field="Appointment Date and Time",
        )
        == "09/06/2026 02:05 PM"
    )


def test_format_encounter_date_value_handles_blank_optional_dates() -> None:
    assert format_encounter_date_value(None, section="Applicant Details", field="Date of Birth") is None
    assert format_encounter_date_value("", section="Applicant Details", field="Date of Birth") is None


def test_format_encounter_date_value_rejects_unrecognized_date_text() -> None:
    try:
        format_encounter_date_value("09.06.2026", section="Applicant Details", field="Date of Birth")
    except ValueError as exc:
        assert "MM/DD/YYYY" in str(exc)
    else:
        raise AssertionError("Expected invalid Encounter date text to raise ValueError.")


def test_create_encounter_workbook_with_interpreter_details(tmp_path) -> None:
    workbook_path = tmp_path / "Encounter_Test_Data.xlsx"

    create_encounter_test_data_workbook(
        workbook_path, include_interpreter_details=True,
    )
    validate_encounter_workbook(workbook_path, include_interpreter_details=True)

    workbook = load_workbook(workbook_path, data_only=True)
    worksheet = workbook[ENCOUNTER_SHEET_NAME]

    actual_columns = []
    current_section = None
    for column_index in range(1, worksheet.max_column + 1):
        section = worksheet.cell(row=1, column=column_index).value
        field = worksheet.cell(row=2, column=column_index).value
        current_section = section or current_section
        actual_columns.append((current_section, field))

    assert worksheet.max_column == 87
    assert actual_columns == encounter_columns(include_interpreter_details=True)

    section_names = []
    for sec, _ in actual_columns:
        if not section_names or section_names[-1] != sec:
            section_names.append(sec)

    ud_idx = section_names.index("Upload Documents")
    id_idx = section_names.index("Interpreter Details")
    assert id_idx == ud_idx + 1

    interpreter_fields = [
        field for sec, field in actual_columns if sec == "Interpreter Details"
    ]
    assert interpreter_fields == list(INTERPRETER_DETAILS_FIELDS)

    encounter_data = read_encounter_test_data(workbook_path=workbook_path)
    assert "Interpreter Details" in encounter_data
    assert list(encounter_data["Interpreter Details"]) == list(INTERPRETER_DETAILS_FIELDS)

    workbook.close()
