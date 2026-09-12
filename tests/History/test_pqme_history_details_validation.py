import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from typing import Any

from openpyxl import load_workbook
from playwright.sync_api import Locator, Page, expect  # pyrefly: ignore[missing-import]

from conftest import test_case as report_case
from utils.auth import authenticate_super_admin
from utils.business_days import calculate_report_due_date
from utils.encounter_test_data import (
    PQME_ENCOUNTER_SHEET_NAME,
    PQME_ENCOUNTER_WORKBOOK_PATH,
    read_encounter_test_data,
)


REQUIRED_EXCEL_FIELDS = (
    ("Applicant Details", "First Name"),
    ("Applicant Details", "Last Name"),
    ("Applicant Details", "Email"),
    ("Encounter Details", "Physician"),
    ("Encounter Details", "Appointment Date and Time"),
)


class HistoryValidationError(AssertionError):
    """Raised when preconditions or History Details validation fail."""


@dataclass(frozen=True)
class MatchedEncounter:
    row: Locator
    applicant_table_name: str
    physician_table_name: str
    evaluation_date: date


def _visible(locator: Locator) -> Locator:
    return locator.first


def _normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _normalize_for_compare(value: str) -> str:
    return _normalize_text(value).casefold()


def _required_value(data: dict[str, dict[str, Any]], section: str, field: str, row_number: int) -> str:
    value = data.get(section, {}).get(field)
    if value is None or not str(value).strip():
        raise HistoryValidationError(
            f"PQME Excel row {row_number} is missing required field: {section} -> {field}"
        )
    return str(value).strip()


def _first_actual_encounter_row() -> int:
    workbook = load_workbook(PQME_ENCOUNTER_WORKBOOK_PATH, read_only=True, data_only=True)
    try:
        if PQME_ENCOUNTER_SHEET_NAME not in workbook.sheetnames:
            raise HistoryValidationError(
                f"Worksheet '{PQME_ENCOUNTER_SHEET_NAME}' not found in {PQME_ENCOUNTER_WORKBOOK_PATH}."
            )
        worksheet = workbook[PQME_ENCOUNTER_SHEET_NAME]
        for row_number in range(3, worksheet.max_row + 1):
            data = read_encounter_test_data(
                row_number=row_number,
                workbook_path=PQME_ENCOUNTER_WORKBOOK_PATH,
                sheet_name=PQME_ENCOUNTER_SHEET_NAME,
            )
            if all(str(data.get(section, {}).get(field) or "").strip() for section, field in REQUIRED_EXCEL_FIELDS):
                return row_number
    finally:
        workbook.close()

    raise HistoryValidationError(
        f"No populated PQME encounter data row found in {PQME_ENCOUNTER_WORKBOOK_PATH} "
        f"worksheet '{PQME_ENCOUNTER_SHEET_NAME}'."
    )


def _applicant_name(data: dict[str, dict[str, Any]]) -> str:
    applicant = data["Applicant Details"]
    return " ".join(
        part
        for part in (
            str(applicant.get("First Name") or "").strip(),
            str(applicant.get("Middle Name") or "").strip(),
            str(applicant.get("Last Name") or "").strip(),
        )
        if part
    )


def _applicant_table_name(data: dict[str, dict[str, Any]]) -> str:
    applicant = data["Applicant Details"]
    first_name = str(applicant.get("First Name") or "").strip()
    middle_name = str(applicant.get("Middle Name") or "").strip()
    last_name = str(applicant.get("Last Name") or "").strip()
    given_names = " ".join(part for part in (first_name, middle_name) if part)
    return f"{last_name}, {given_names}" if last_name and given_names else _applicant_name(data)


def _physician_table_name(value: str) -> str:
    parts = [part.strip() for part in value.split(",") if part.strip()]
    if len(parts) >= 2:
        first_part = re.sub(r"\b(?:MD|DO)\b\.?", "", parts[0], flags=re.I).strip()
        second_part = re.sub(r"\b(?:MD|DO)\b\.?", "", parts[1], flags=re.I).strip()
        if first_part.casefold() == "physician":
            return second_part
        return " ".join(part for part in (second_part, first_part) if part)
    without_credentials = re.sub(r"\b(?:MD|DO)\b\.?", "", value, flags=re.I)
    without_generic_title = re.sub(r"(^|\s)Physician($|\s)", " ", without_credentials, flags=re.I)
    return _normalize_text(without_generic_title)


def _format_ui_date(value: str) -> str:
    digits = re.findall(r"\d+", value)
    if len(digits) >= 3:
        month, day, year = digits[0], digits[1], digits[2]
        return f"{int(month):02d}/{int(day):02d}/{int(year):04d}"
    return value.strip()


def _parse_date(value: str) -> date:
    normalized = _format_ui_date(value)
    return datetime.strptime(normalized, "%m/%d/%Y").date()


def calculate_expected_history_due_date(evaluation_date: date) -> date:
    return evaluation_date - timedelta(days=3)


def _contains_display_text(actual_text: str, expected_text: str) -> bool:
    return _normalize_for_compare(expected_text) in _normalize_for_compare(actual_text)


def _open_view_encounters(page: Page) -> None:
    view_encounters = page.get_by_role("link", name=re.compile(r"View Encounters", re.I)).or_(
        page.get_by_text("View Encounters", exact=True)
    )
    expect(_visible(view_encounters)).to_be_visible()
    _visible(view_encounters).click()
    page.wait_for_url(re.compile(r".*/#/encounter_details$"))
    expect(page.get_by_role("button", name=re.compile(r"New encounter", re.I))).to_be_visible()


def _find_matching_encounter_row(page: Page, data: dict[str, dict[str, Any]], row_number: int) -> MatchedEncounter:
    applicant_table_name = _applicant_table_name(data)
    physician = _required_value(data, "Encounter Details", "Physician", row_number)
    physician_table_name = _physician_table_name(physician)
    appointment_value = _required_value(data, "Encounter Details", "Appointment Date and Time", row_number)
    evaluation_date = _parse_date(appointment_value)
    evaluation_date_text = evaluation_date.strftime("%m/%d/%Y")

    search = page.get_by_placeholder(re.compile("Applicant name", re.I)).or_(
        page.get_by_role("textbox", name=re.compile("Applicant name", re.I))
    )
    if search.count():
        search.first.fill(applicant_table_name)

    row = (
        page.get_by_role("row")
        .filter(has_text=applicant_table_name)
        .filter(has_text="PQME")
        .filter(has_text=evaluation_date_text)
    )
    expect(
        _visible(row),
        (
            f"Precondition failed: PQME encounter from first Excel data row {row_number} "
            f"was not found in View Encounters. Expected Applicant '{applicant_table_name}', "
            f"Physician '{physician_table_name}', Evaluation Date {evaluation_date_text}."
        ),
    ).to_be_visible(timeout=30000)

    matched_row = _visible(row)
    row_text = matched_row.inner_text()
    assert _contains_display_text(row_text, applicant_table_name), (
        f"Applicant Name mismatch in View Encounters. Expected '{applicant_table_name}', "
        f"found row text {row_text!r}."
    )
    assert _contains_display_text(row_text, physician_table_name), (
        f"Physician mismatch in View Encounters. Expected '{physician_table_name}', "
        f"found row text {row_text!r}."
    )
    assert evaluation_date_text in row_text, (
        f"Date of Evaluation mismatch in View Encounters. Expected {evaluation_date_text}, "
        f"found row text {row_text!r}."
    )

    return MatchedEncounter(
        row=matched_row,
        applicant_table_name=applicant_table_name,
        physician_table_name=physician_table_name,
        evaluation_date=evaluation_date,
    )


def _open_matched_encounter(matched: MatchedEncounter) -> None:
    applicant_link = matched.row.get_by_role("link", name=re.compile(re.escape(matched.applicant_table_name), re.I))
    if applicant_link.count() == 0:
        applicant_link = matched.row.get_by_text(matched.applicant_table_name, exact=True)
    expect(applicant_link.first).to_be_visible()
    applicant_link.first.click()


def _open_history_tab(page: Page) -> None:
    history_tab = page.get_by_role("tab", name=re.compile(r"^\s*History\s*$", re.I)).or_(
        page.get_by_text(re.compile(r"^\s*History\s*$", re.I))
    )
    expect(history_tab.first, "History tab is missing or not visible for the selected PQME encounter.").to_be_visible(
        timeout=30000
    )
    history_tab.first.click()
    expect(page.get_by_text(re.compile(r"^\s*History Details\s*$", re.I)).first).to_be_visible(timeout=30000)


def _history_details_section(page: Page) -> Locator:
    heading = page.get_by_text(re.compile(r"^\s*History Details\s*$", re.I)).first
    expect(heading).to_be_visible()
    main = page.locator("main").first
    expect(main).to_contain_text(re.compile(r"History details", re.I))
    return main


def _history_field_text(section: Locator, label: str) -> str:
    section_text = _normalize_text(section.inner_text(timeout=5000))
    parsed_value = _read_history_label_value(section_text, label)
    if parsed_value:
        return parsed_value

    label_pattern = re.compile(rf"^\s*{re.escape(label)}\s*(?:\*)?\s*$", re.I)
    labeled_form_field = section.locator("mat-form-field").filter(has_text=label_pattern)
    if labeled_form_field.count():
        value_text = labeled_form_field.first.locator("input, textarea, [role='textbox'], mat-select").first
        if value_text.count():
            value = value_text.input_value() if value_text.evaluate("el => 'value' in el") else value_text.inner_text()
            if _normalize_text(value):
                return _normalize_text(value)
        text_without_label = re.sub(label_pattern, "", labeled_form_field.first.inner_text()).strip()
        if text_without_label:
            return _normalize_text(text_without_label)

    label_node = section.get_by_text(label_pattern).first
    expect(label_node, f"History Details field label not found: {label}").to_be_visible()
    container = label_node.locator("xpath=ancestor::*[self::div or self::li or self::p][1]")
    text = container.inner_text()
    value = re.sub(label_pattern, "", text).strip()
    if value:
        return _normalize_text(value)

    following = label_node.locator("xpath=following::*[normalize-space()][1]")
    expect(following).to_be_visible()
    return _normalize_text(following.inner_text())


def _read_history_label_value(section_text: str, label: str) -> str:
    label_aliases = {
        "Physician": ("Physician",),
        "Date of Evaluation": ("Date of evaluation", "Date of Evaluation"),
        "Report Due Date": ("Report due date", "Report Due Date"),
        "History Due Date": ("History due date", "History Due Date"),
        "Applicant Email Address": ("Applicant email", "Applicant Email Address", "Applicant Email"),
    }
    labels_in_order = (
        "Physician",
        "Applicant",
        "Date of evaluation",
        "Report due date",
        "History due date",
        "Applicant email",
        "Historian",
        "Last edited on",
        "Contact applicant",
    )

    aliases = label_aliases.get(label, (label,))
    alias_pattern = "|".join(re.escape(alias) for alias in aliases)
    boundary_pattern = "|".join(re.escape(item) for item in labels_in_order)
    match = re.search(
        rf"(?:^|\s)(?:{alias_pattern})\s+(.*?)(?=\s+(?:{boundary_pattern})\b|$)",
        section_text,
        flags=re.I,
    )
    return _normalize_text(match.group(1)) if match else ""


def _assert_same_text(field_name: str, expected: str, actual: str) -> None:
    assert _normalize_for_compare(actual) == _normalize_for_compare(expected), (
        f"{field_name} mismatch. Expected '{expected}', found '{actual}'."
    )


def _assert_same_date(field_name: str, expected: date, actual_text: str) -> None:
    actual = _parse_date(actual_text)
    assert actual == expected, (
        f"{field_name} mismatch. Expected {expected.strftime('%m/%d/%Y')}, "
        f"found {actual.strftime('%m/%d/%Y')}."
    )


@report_case(
    test_case_id="SUPERADMIN_PQME_HISTORY_TC_001",
    module="PQME - History",
    test_case_name="PQME History Details Validation",
    title=(
        "Verify that the History tab of an existing PQME encounter displays the correct "
        "Physician, Evaluation Date, Report Due Date, History Due Date, and Applicant "
        "Email Address."
    ),
    test_type="Positive",
    preconditions=(
        "A PQME encounter matching the first encounter data row in the PQME Excel "
        "test-data workbook already exists."
    ),
    steps=(
        "1. Read the first actual PQME encounter data row from the PQME Excel workbook.\n"
        "2. Login as Super Admin using existing environment credentials and verification helper.\n"
        "3. Navigate to View Encounters.\n"
        "4. Search for the PQME encounter by Applicant name.\n"
        "5. Identify the exact row using Applicant name, Physician, Encounter Type, and Date of Evaluation.\n"
        "6. Open the encounter by clicking the Applicant name from the matched row.\n"
        "7. Verify the History tab is visible and open it.\n"
        "8. Validate History Details in order: Physician, Date of Evaluation, Report Due Date, "
        "History Due Date, Applicant Email Address."
    ),
    test_data=(
        "First actual PQME Encounter data row from test_data/PQME_Encounter_Test_Data.xlsx, "
        "worksheet Encounter_Test_Data. Super Admin password and verification PIN are not reported."
    ),
    expected_result=(
        "The correct PQME encounter should be found in View Encounters, the History tab should "
        "be visible, and History Details should show the correct Physician, Evaluation Date "
        "matching Appointment Date and Time, Report Due Date equal to Evaluation Date + 30 "
        "calendar days adjusted backward for weekends and U.S. federal holidays, History Due "
        "Date equal to Evaluation Date - 3 calendar days, and Applicant Email matching the "
        "PQME Excel data."
    ),
    actual_result=(
        "The existing PQME encounter was found, opened from View Encounters, the History tab "
        "was opened, and the requested History Details fields were validated in order."
    ),
)
def test_validate_pqme_history_details(page: Page) -> None:
    stage = "reading first PQME encounter Excel data row"
    try:
        row_number = _first_actual_encounter_row()
        encounter_data = read_encounter_test_data(
            row_number=row_number,
            workbook_path=PQME_ENCOUNTER_WORKBOOK_PATH,
            sheet_name=PQME_ENCOUNTER_SHEET_NAME,
        )
        for section, field in REQUIRED_EXCEL_FIELDS:
            _required_value(encounter_data, section, field, row_number)

        applicant_name = _applicant_name(encounter_data)
        applicant_email = _required_value(encounter_data, "Applicant Details", "Email", row_number)
        physician = _physician_table_name(_required_value(encounter_data, "Encounter Details", "Physician", row_number))

        stage = "logging in as Super Admin"
        authenticate_super_admin(page)

        stage = "opening View Encounters"
        _open_view_encounters(page)

        stage = "finding matching PQME encounter row"
        matched = _find_matching_encounter_row(page, encounter_data, row_number)

        stage = "opening the matched PQME encounter"
        _open_matched_encounter(matched)

        stage = "opening History tab"
        _open_history_tab(page)
        history_details = _history_details_section(page)

        evaluation_date = matched.evaluation_date
        report_due_date = calculate_report_due_date(evaluation_date)
        history_due_date = calculate_expected_history_due_date(evaluation_date)
        base_plus_30 = evaluation_date + timedelta(days=30)

        stage = "validating History Details Physician"
        actual_physician = _history_field_text(history_details, "Physician")
        _assert_same_text("History Physician", physician, actual_physician)

        stage = "validating History Details Date of Evaluation"
        actual_evaluation_date = _history_field_text(history_details, "Date of Evaluation")
        _assert_same_date("History Date of Evaluation", evaluation_date, actual_evaluation_date)

        stage = "validating History Details Report Due Date"
        actual_report_due_date = _history_field_text(history_details, "Report Due Date")
        actual_report_due_date_obj = _parse_date(actual_report_due_date)
        assert actual_report_due_date_obj == report_due_date, (
            f"History Report Due Date mismatch. "
            f"Evaluation Date: {evaluation_date.strftime('%m/%d/%Y')}, "
            f"Base +30 Date: {base_plus_30.strftime('%m/%d/%Y')} ({base_plus_30.strftime('%A')}), "
            f"Expected adjusted Report Due Date: {report_due_date.strftime('%m/%d/%Y')}, "
            f"Actual History Report Due Date: {actual_report_due_date_obj.strftime('%m/%d/%Y')}."
        )

        stage = "validating History Details History Due Date"
        actual_history_due_date = _history_field_text(history_details, "History Due Date")
        actual_history_due_date_obj = _parse_date(actual_history_due_date)
        assert actual_history_due_date_obj == history_due_date, (
            f"History Due Date mismatch. "
            f"Evaluation Date: {evaluation_date.strftime('%m/%d/%Y')}, "
            f"Expected History Due Date: {history_due_date.strftime('%m/%d/%Y')}, "
            f"Actual History Due Date: {actual_history_due_date_obj.strftime('%m/%d/%Y')}."
        )

        stage = "validating History Details Applicant Email Address"
        actual_email = _history_field_text(history_details, "Applicant Email Address")
        _assert_same_text("History Applicant Email Address", applicant_email, actual_email)

        print(
            "PQME History Details validated: "
            f"Excel Row={row_number}; Applicant={applicant_name}; "
            f"View Applicant={matched.applicant_table_name}; "
            f"Physician={actual_physician}; "
            f"Date of Evaluation={actual_evaluation_date}; "
            f"Report Due Date={actual_report_due_date}; "
            f"History Due Date={actual_history_due_date}; "
            f"Applicant Email={actual_email}"
        )
    except HistoryValidationError:
        raise
    except AssertionError:
        raise
    except Exception as exc:
        raise AssertionError(
            f"PQME History Details Validation failed while {stage} ({type(exc).__name__})."
        ) from None
